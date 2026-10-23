"""Local Junbin/PoLARIS test of Dan's existing GP and gate APIs; no LLM or hardware.

This is a bounded, hand-annotated research probe, not a complete paper extraction
or a FlowLab Platform. Source data stays outside Git. See assessment.md.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
CSV_SHA = '113aa4853fb7a408a8707c6f73853b2db3754c2829779486c58eb7a967bf1dca'
XLSX_SHA = 'fa5b0e7cae80eba9669f1dd96434e98562237d23d30610ed9001aae2abe3a205'
POINT = {'Cs': 57.2091, 'Ag': 60., 'Na': 25., 'In-HCl': 39.05705,
         'Bi': 25., 'ODE': 60., 'Trxn': 220.}


def tag(value, provenance, source, note=None):
    return dict(value=value, provenance=provenance, source=source, note=note)


def stated(value, source):
    return tag(value, 'paper_stated', source)


def assumed(value, note):
    return tag(value, 'assumed', 'Local reconstruction policy', note)


def make_probe_card(csv):
    from xp_plus.schema import ModelCard
    inputs = []
    for name, value in POINT.items():
        temp = name == 'Trxn'
        inputs.append(dict(name=name, kind='constant', column=name,
            definition=stated('Reaction temperature' if temp else f'Absolute {name} feed flow rate',
                             'Junbin.pdf pp4-5; deposited Digital Twin.py input_indices'),
            value=tag(value, 'inferred', 'JunbinSupplemental.pdf p3 Note 1; SourceData.xlsx Figure 3b A3:B9',
                      'Inverse SI normalization: F=25+35*x, T=170+50*x. Fixed-condition probe only; no dynamic pump binding.'),
            unit='degC' if temp else 'uL/min',
            low=stated(170. if temp else 25., 'JunbinSupplemental.pdf p3 Note 1'),
            high=stated(220. if temp else 60., 'JunbinSupplemental.pdf p3 Note 1')))
    return ModelCard.model_validate(dict(id='polaris_gp', family='gaussian_process',
        purpose='Figure 4a CV and Figure 3b fixed-condition computational probe; not mechanistic reproduction',
        inputs=inputs, outputs=[dict(name='plqy_proxy', column='Unnamed: 8', unit='fraction',
            definition=stated('PLQY proxy in the deposited initialization CSV last column', 'Junbin.pdf p5; deposited Digital Twin.py'),
            low=assumed(0., 'Identity output mapping for deposited fractional proxy; CV uses CSV targets directly.'),
            high=assumed(1., 'Identity output mapping for deposited fractional proxy; CV uses CSV targets directly.'))],
        library=stated('BoTorch SingleTaskGP', 'Zenodo 19665377 Digital Twin.py'),
        code=stated('https://zenodo.org/records/19665377', 'Junbin.pdf p10 ref53'),
        kernel=assumed('Installed SingleTaskGP default: ARD RBF',
                       'Current-backend probe only. Original kernel/version not pinned; legacy is a separate sensitivity probe.'),
        hyperparameter_fitting=stated('ExactMarginalLogLikelihood optimized with fit_gpytorch_mll', 'Deposited Digital Twin.py'),
        training_data=tag(str(csv.resolve()), 'inferred', 'Zenodo 19665377, Digital Twin/Initialization_Campaign.csv',
                          'Local copy verified by SHA256; this path is not paper metadata.'),
        training_data_sha256=CSV_SHA, training_data_normalized=True,
        n_training_points=stated(80., 'Junbin.pdf p5'),
        validation=stated('10-fold CV; reported R2 0.71', 'Junbin.pdf p5 Figure 4a'),
        validation_protocol=dict(folds=10, pre_shuffle_seed=None, split_seed=42, metric='r2'),
        metric_definitions=stated('sklearn.metrics.r2_score', 'Deposited Digital Twin.py cross_val_parity_plot'),
        seeds=stated('KFold shuffle=True random_state=42; no GP fitting seed specified', 'Deposited Digital Twin.py')))


def workbook_metrics(path):
    import numpy as np
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        pairs = np.array(list(wb['Figure 4a'].iter_rows(min_row=3, max_row=82, min_col=1, max_col=2, values_only=True)), float)
        x, y = np.array(list(wb['Supplementary Figure 6'].iter_rows(min_row=3, max_row=8, min_col=1, max_col=2, values_only=True)), float).T
        result = {'sheets': len(wb.sheetnames), 'parity_range': 'Figure 4a!A3:B82',
                  'parity_r2': float(1 - ((pairs[:, 0] - pairs[:, 1])**2).sum() / ((pairs[:, 0] - pairs[:, 0].mean())**2).sum()),
                  'calibration_range': 'Supplementary Figure 6!A3:B8',
                  'zero_intercept_calibration_slope': float(x @ y / (x @ x)), 'stability': {}}
        for sheet, last in [('Figure 2c', 1182), ('Figure 3c', 748), ('Supplementary Figure 2', 1182), ('Supplementary Figure 3', 1182)]:
            values = [r[0] for r in wb[sheet].iter_rows(min_row=3, max_row=last, min_col=2, max_col=2, values_only=True) if isinstance(r[0], (int, float))]
            a = np.array(values)
            result['stability'][sheet] = dict(range=f'B3:B{last}', n=len(a), population_cv_pct=float(a.std(ddof=0)/a.mean()*100))
        return result, pairs
    finally:
        wb.close()


def gate_check(card, out):
    from pydantic import ValidationError
    from xp_plus.extraction import Extraction
    from xp_plus.pipeline import run_pipeline
    from xp_plus.resolve import apply_resolutions
    from xp_plus.schema import Component, Connection, Rig, PaperRef, PaperFeatures, Claim, Gap, ModelInput, Resolution
    from xp_plus.twin import run_gate
    from xp_plus.twin.gate import check_model
    card = card.model_copy(deep=True)
    resolutions = []
    for obj, field, target in [(card, 'kernel', 'kernel'), (card.outputs[0], 'low', 'outputs.plqy_proxy.low'), (card.outputs[0], 'high', 'outputs.plqy_proxy.high')]:
        v = getattr(obj, field)
        resolutions.append(Resolution(target='models.polaris_gp.'+target, value=v.value, provenance=v.provenance, source=v.source, note=v.note))
        setattr(obj, field, type(v)())
    components, connections = [], []
    for i in card.inputs[:-1]:
        pump, res = 'pump_'+i.name, 'res_'+i.name
        components += [Component(id=res, kind='reservoir', name=i.name), Component(id=pump, kind='pump',
            manufacturer=stated('Chemyx', 'Junbin.pdf p8 Hardware'), model=stated('Fusion 6000', 'Junbin.pdf p8 Hardware'), flow_rate_uL_min=i.value)]
        connections += [Connection(src=res, dst=pump, direct=True), Connection(src=pump, dst='manifold')]
    components += [Component(id='res_pfo', kind='reservoir'), Component(id='pump_pfo', kind='pump', manufacturer=stated('Chemyx', 'Junbin.pdf p8 Hardware'), model=stated('Fusion 6000', 'Junbin.pdf p8 Hardware')),
        Component(id='manifold', kind='mixer', name='P-150 seven-port manifold'), Component(id='segmenting_tee', kind='mixer', name='P-714 carrier tee'),
        Component(id='reactor', kind='reactor', temperature_C=card.inputs[-1].value), Component(id='cooling', kind='tubing', name='Ambient cooling section'),
        Component(id='optical_group', kind='detector', name='Abstract group of two actual optical cells; order unresolved')]
    connections += [Connection(src=a, dst=b, direct=d) for a,b,d in [('res_pfo','pump_pfo',True), ('pump_pfo','segmenting_tee',False), ('manifold','segmenting_tee',False), ('segmenting_tee','reactor',False), ('reactor','cooling',False), ('cooling','optical_group',False)]]
    f = PaperFeatures(paper=PaperRef(title='Autonomous microfluidic experimentation for exploring reaction inference and synthesizing double perovskite nanoplatelets', doi='10.1038/s41467-026-72765-2', year=2026),
        experiment='Partial unit-operation gate probe of Figure 4a/3b; not a complete extraction or FlowLab Platform', rig=Rig(components=components, connections=connections), models=[card],
        claims=[Claim(id='cv', metric='cv_r2_plqy_proxy', value=.71, tolerance=.005, tolerance_basis='Local two-decimal rounding criterion, not measured uncertainty', source='Junbin.pdf p5 Figure 4a')])
    gaps = [Gap(id='extraction:'+key, target=target, description=text, detected_by='extraction') for key,target,text in [
        ('absolute_flow_binding','models.polaris_gp.inputs','Fixed inputs cannot represent dynamic absolute pump flows; no flow_rate feature exists.'),
        ('segmented_physics','rig','PFO/reactive-droplet segmentation is not represented by the single-phase transport. A GP proxy is not a mechanistic validation.'),
        ('optical_arrangement','rig.components.optical_group','Two actual cells aggregated for this partial test; individual geometry, order and coupling remain unestablished from reviewed evidence.')]]
    extraction = Extraction(features=f, gaps=gaps)
    gate = run_gate(f)
    reviewed = apply_resolutions(f, resolutions)
    standard = run_pipeline(extraction, resolutions=resolutions)
    assert not gate.passed and standard.verdict == 'insufficient_data'
    assert standard.twin is None and standard.physical is None and not standard.computational
    assert not any(g.blocking for g in check_model(reviewed))
    try:
        ModelInput(name='Cs', kind='flow_rate')
    except ValidationError:
        rejected = True
    else:
        rejected = False
    assert rejected
    (out/'original-extraction.json').write_text(extraction.model_dump_json(indent=2), encoding='utf-8')
    (out/'resolutions.json').write_text(json.dumps([r.model_dump() for r in resolutions], indent=2), encoding='utf-8')
    standard.save(out/'blocked-standard.json')
    return {'scope': f.experiment, 'original_blockers': [g.target for g in gate.gaps if g.blocking],
            'model_gate_after_explicit_defaults': 'passed', 'pipeline_verdict': standard.verdict,
            'pipeline_computational_records': len(standard.computational),
            'remaining_blockers': [g.target for g in standard.gaps if g.blocking and g.status=='open'],
            'absolute_flow_feature_rejected': rejected}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--csv', required=True, type=Path, help='Verified Digital Twin/Initialization_Campaign.csv from Zenodo 19665377')
    p.add_argument('--workbook', required=True, type=Path, help='User-supplied SourceData.xlsx')
    p.add_argument('--out', type=Path, default=ROOT/'.cache/xp_plus/polaris-check')
    args = p.parse_args()
    for path, expected in [(args.csv, CSV_SHA), (args.workbook, XLSX_SHA)]:
        if hashlib.sha256(path.read_bytes()).hexdigest() != expected:
            p.error(f'Source identity mismatch: {path}. Review a new version before changing this fixture.')
    args.out.mkdir(parents=True, exist_ok=True)
    import numpy as np
    import torch
    from xp_plus.blackbox.validate import cross_validate
    from xp_plus.blackbox.model import BlackBoxModel, BACKENDS
    torch.set_num_threads(1)
    card = make_probe_card(args.csv)
    (args.out/'probe-model-card.json').write_text(card.model_dump_json(indent=2), encoding='utf-8')
    wb, pairs = workbook_metrics(args.workbook)
    result = {'csv_sha256': CSV_SHA, 'workbook_sha256': XLSX_SHA, 'workbook': wb, 'gate': gate_check(card, args.out), 'cv': {}}
    for backend in ('botorch', 'botorch_legacy'):
        print(f'Running {backend}: 80 rows, seven inputs, 10 folds, split seed 42', flush=True)
        t0 = time.monotonic()
        pred, act = cross_validate(card, backend)
        r2 = float(1 - ((act-pred)**2).sum()/((act-act.mean())**2).sum())
        result['cv'][backend] = {'r2': r2, 'backend': BACKENDS[backend].version(), 'elapsed_s': time.monotonic()-t0,
            'observations_match_workbook_at_5_decimals': bool(np.allclose(np.round(act[:,0],5), pairs[:,0], rtol=0, atol=1e-12)),
            'prediction_rmse_vs_workbook': float(np.sqrt(np.mean((pred[:,0]-pairs[:,1])**2))),
            'reported_0_71_rounding_match': bool(abs(r2-.71) <= .005)}
        assert result['cv'][backend]['observations_match_workbook_at_5_decimals']
        np.savez(args.out/f'{backend}-cv.npz', predicted=pred, actual=act)
        print(json.dumps(result['cv'][backend]), flush=True)
    m = BlackBoxModel(card, 'botorch').fit()
    mean, std = m.predict(np.array([list(POINT.values())]))
    result['fixed_point_probe'] = {'inputs': POINT, 'mean_proxy_fraction': float(mean[0,0]), 'posterior_std_fraction': float(std[0,0]), 'warning': 'Computational prediction only; not measured PLQY, hardware evidence, or segmented-flow validation.'}
    (args.out/'results.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(f'Saved results to {args.out}', flush=True)


if __name__ == '__main__':
    main()
